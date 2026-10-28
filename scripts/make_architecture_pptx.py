#!/usr/bin/env python3
"""Generate a NATIVE, editable PowerPoint architecture slide.

Mirrors docs/architecture.html as real PowerPoint shapes (bands, icon pictures,
text boxes, connectors, legend, metrics) so judges/teammates can edit every
element in PowerPoint — not a flat screenshot.

Icons are the official AWS Architecture Icons, rasterized to PNG under
docs/assets/icons/pptx_png.

Usage:  python3 scripts/make_architecture_pptx.py
Output: docs/One-Touch-Resolve-Architecture.pptx
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ICONS = ROOT / "docs" / "assets" / "icons" / "pptx_png"
OUT = ROOT / "docs" / "One-Touch-Resolve-Architecture.pptx"

EMU_W, EMU_H = 12192000, 6858000  # 13.333in x 7.5in (16:9)
GX, GY = 1600, 900                # virtual grid
SX, SY = EMU_W / GX, EMU_H / GY


def X(v): return Emu(int(v * SX))
def Y(v): return Emu(int(v * SY))


INK = RGBColor(0x23, 0x2F, 0x3E)
INK2 = RGBColor(0x5A, 0x69, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLUE = RGBColor(0x00, 0x73, 0xBB)
PURPLE = RGBColor(0x7D, 0x3C, 0xCF)
TEAL = RGBColor(0x01, 0xA8, 0x8D)
ORANGE = RGBColor(0xED, 0x71, 0x00)
VIOLET = RGBColor(0x5F, 0x1E, 0xBE)
GREY = RGBColor(0x87, 0x95, 0xA4)

BANDS = [
    ("1 · AWS Communication Developer Services — communications backbone", 60, 120, 1480, 150, BLUE, RGBColor(0xF6, 0xFB, 0xFF)),
    ("2 · Ingest & Dispatch", 60, 290, 1480, 150, PURPLE, RGBColor(0xFB, 0xF9, 0xFF)),
    ("3 · Multi-Agent Orchestration — Amazon Bedrock AgentCore (Runtime · Memory · Strands · Nova)", 60, 460, 1480, 210, TEAL, RGBColor(0xF5, 0xFC, 0xFA)),
    ("4 · State, Persistence & Observability", 60, 690, 1480, 120, ORANGE, RGBColor(0xFF, 0xFA, 0xF4)),
]

# (icon, title, sub, svc, x, y)
NODES = [
    ("whatsapp", "EUM Social", "WhatsApp — FNOL media", "AWS End User Messaging Social", 110, 150),
    ("rcs", "EUM RCS / SMS", "cards · RCS->SMS fallback", "AWS End User Messaging", 380, 150),
    ("ses", "Amazon SES", "coverage / settlement", "Amazon SES", 660, 150),
    ("router", "Channel router", "right channel per party", "", 1120, 150),
    ("sns", "Amazon SNS", "topic per channel", "", 230, 320),
    ("lambda", "Dispatcher", "normalize-dedupe-invoke", "", 470, 320),
    ("s3", "Amazon S3", "FNOL photos/video", "", 730, 320),
    ("dynamodb", "Dedupe table", "idempotency · TTL", "", 980, 320),
    ("agentcore", "Supervisor", "plans-reconciles-re-plans", "", 200, 500),
    ("claimant", "Sub-agents x4", "Claimant-Repair-Parts-Lien", "", 1230, 505),
    ("lambda", "Tools (Lambda)", "open-estimate-settle", "", 1390, 505),
    ("agentcore", "AgentCore Memory", "state by actor_id", "", 1250, 430),
    ("bedrock", "Bedrock · Nova", "triage + tool-calling", "", 1470, 640),
    ("dynamodb", "coe-claims", "source of truth · PITR", "", 250, 700),
    ("cloudwatch", "CloudWatch", "DLQ + error alarms", "", 560, 700),
    ("iam", "IAM least-priv", "scoped send-ddb-bedrock", "", 830, 700),
    ("dynamodb", "coe-fleet", "managed rental fleet", "", 1100, 700),
]

# (x1,y1,x2,y2,color,dashed,label)
WIRES = [
    (300, 270, 300, 320, GREY, False, ""),
    (340, 360, 455, 360, GREY, False, ""),
    (585, 360, 725, 360, GREY, False, "media"),
    (520, 380, 300, 500, GREY, False, ""),
    (330, 560, 455, 560, TEAL, False, "A2A"),
    (1160, 560, 1225, 560, TEAL, False, ""),
    (1360, 560, 1388, 560, TEAL, False, ""),
    (1440, 580, 1520, 640, ORANGE, False, "InvokeModel"),
    (775, 360, 1545, 650, ORANGE, True, "FNOL media -> Nova vision"),
    (1420, 580, 320, 700, PURPLE, True, "persist"),
    (1430, 580, 1165, 700, PURPLE, False, ""),
    (1150, 540, 1480, 270, PURPLE, True, "outbound CDS"),
]

LEGEND = [
    (GREY, "Inbound / control"),
    (TEAL, "Agent (A2A) calls"),
    (ORANGE, "Bedrock InvokeModel"),
    (PURPLE, "Outbound CDS / persist"),
]

METRICS = [
    ("3 CDS", "WhatsApp-RCS/SMS-SES"),
    ("1+4", "supervisor + sub-agents"),
    ("1 gate", "commit + human approval"),
    ("1 CDK", "one-command deploy"),
    ("re-skin", "swap roster -> vertical"),
]


def _fill(shape, rgb):
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb


def _text(tb, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(20000)
    tf.margin_top = tf.margin_bottom = Emu(10000)
    first = True
    for text, size, color, bold in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = text
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
        r.font.name = "Arial"
        first = False


def _dash(shape):
    ln = shape.line._get_or_add_ln()
    d = ln.find(qn("a:prstDash"))
    if d is None:
        d = ln.makeelement(qn("a:prstDash"), {}); ln.append(d)
    d.set("val", "dash")


def _arrow(conn):
    ln = conn.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"),
                             {"type": "triangle", "w": "med", "len": "med"}))


def build():
    prs = Presentation()
    prs.slide_width = Emu(EMU_W); prs.slide_height = Emu(EMU_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes

    t = shapes.add_textbox(X(40), Y(18), X(1000), Y(60))
    _text(t, [("One Touch Resolve - AWS Architecture", 26, INK, True)])
    s = shapes.add_textbox(X(41), Y(66), X(1100), Y(30))
    _text(s, [("Multi-party claims orchestration · AWS Communication Developer Services + Amazon Bedrock AgentCore", 11, INK2, False)])
    w = shapes.add_textbox(X(1150), Y(20), X(410), Y(50))
    _text(w, [("One Touch Resolve", 12, VIOLET, True),
              ("AWS CDS Agentic AI Partner Hackathon 2026", 8, INK2, False)], align=PP_ALIGN.RIGHT)

    cloud = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(24), Y(108), X(1500), Y(760))
    cloud.fill.background(); cloud.line.color.rgb = INK; cloud.line.width = Pt(1.5)
    cloud.shadow.inherit = False
    ct = shapes.add_textbox(X(40), Y(112), X(500), Y(24))
    _text(ct, [("AWS Cloud · one CDK stack  ·  Region us-east-1", 11, INK, True)])

    for title, x, y, wdt, hgt, accent, bg in BANDS:
        band = shapes.add_shape(MSO_SHAPE.RECTANGLE, X(x), Y(y), X(wdt), Y(hgt))
        _fill(band, bg); band.line.color.rgb = accent; band.line.width = Pt(1.25)
        band.shadow.inherit = False
        bt = shapes.add_textbox(X(x + 12), Y(y + 6), X(wdt - 24), Y(22))
        _text(bt, [(title, 11, accent, True)])

    gate = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(460), Y(520), X(700), Y(130))
    gate.fill.solid(); gate.fill.fore_color.rgb = RGBColor(0xF3, 0xF0, 0xFB)
    gate.line.color.rgb = VIOLET; gate.line.width = Pt(2); gate.shadow.inherit = False
    _dash(gate)
    gt = shapes.add_textbox(X(466), Y(526), X(688), Y(120))
    _text(gt, [("MULTI-PARTY COMMIT GATE · HUMAN APPROVAL", 12, VIOLET, True),
               ("Blocks advance until all required parties COMMIT · re-plans on BLOCKED · then PENDING_APPROVAL: adjuster Approve/Hold - no payout without a human (code-enforced)", 9.5, INK2, False),
               ("CLOSED   ·   RE-PLAN   ·   OPEN   ·   APPROVE", 9, VIOLET, True)],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    for x1, y1, x2, y2, color, dashed, label in WIRES:
        conn = shapes.add_connector(MSO_CONNECTOR.ELBOW, X(x1), Y(y1), X(x2), Y(y2))
        conn.line.color.rgb = color; conn.line.width = Pt(1.75 if dashed else 2.25)
        if dashed:
            _dash(conn)
        _arrow(conn)
        if label:
            lx = (x1 + x2) / 2 - 75; ly = (y1 + y2) / 2 - 11
            lb = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(lx), Y(ly), X(170), Y(22))
            lb.fill.solid(); lb.fill.fore_color.rgb = WHITE
            lb.line.color.rgb = color; lb.line.width = Pt(0.75); lb.shadow.inherit = False
            _text(lb, [(label, 8.5, color, True)], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    for icon, title, sub, svc, x, y in NODES:
        img = ICONS / f"{icon}.png"
        if img.exists():
            shapes.add_picture(str(img), X(x + 34), Y(y), height=Y(52))
        tb = shapes.add_textbox(X(x - 24), Y(y + 56), X(168), Y(62))
        runs = [(title, 11, INK, True), (sub, 9, INK2, False)]
        if svc:
            runs.append((svc, 8, ORANGE, True))
        _text(tb, runs, align=PP_ALIGN.CENTER)

    lg = shapes.add_shape(MSO_SHAPE.RECTANGLE, X(980), Y(760), X(430), Y(92))
    _fill(lg, RGBColor(0xFA, 0xFB, 0xFD)); lg.line.color.rgb = RGBColor(0xD8, 0xE0, 0xE7)
    lg.line.width = Pt(0.75); lg.shadow.inherit = False
    lt = shapes.add_textbox(X(990), Y(763), X(200), Y(18))
    _text(lt, [("ARROW KEY", 9, INK2, True)])
    for i, (color, text) in enumerate(LEGEND):
        col = i % 2; row = i // 2
        lx = 996 + col * 210; ly = 788 + row * 28
        line = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, X(lx), Y(ly + 7), X(lx + 34), Y(ly + 7))
        line.line.color.rgb = color; line.line.width = Pt(2.5); _arrow(line)
        kt = shapes.add_textbox(X(lx + 42), Y(ly), X(170), Y(18))
        _text(kt, [(text, 9.5, INK, False)])

    for i, (v, l) in enumerate(METRICS):
        mx = 60 + i * 180
        m = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(mx), Y(762), X(170), Y(56))
        _fill(m, WHITE); m.line.color.rgb = RGBColor(0xE3, 0xE8, 0xEF); m.line.width = Pt(0.75)
        m.shadow.inherit = False
        _text(m, [(v, 14, VIOLET, True), (l, 8.5, INK2, False)],
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    f = shapes.add_textbox(X(40), Y(872), X(600), Y(20))
    _text(f, [("© 2026 · One Touch Resolve · Confidential", 9, RGBColor(0x8B, 0x93, 0xA7), False)])

    prs.save(str(OUT))
    print("wrote", OUT)


if __name__ == "__main__":
    build()
